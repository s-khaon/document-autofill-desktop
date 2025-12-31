#!/usr/bin/env python3
"""
Word模板填充程序
使用python-docx库，支持{{field_name}}占位符
支持普通段落、表格单元格、页眉和页脚

命令行参数：
1. 模板路径
2. 输出路径 (可选，当操作为extract时不需要)
3. JSON数据 (可选，当操作为extract时不需要)
4. 操作类型 (可选，默认为fill，可取值fill或extract)
"""

import sys
import json
from docx import Document
from docx.document import Document as DocxDocument
from docx.section import _Header, _Footer
from docx.table import Table
from docx.text.paragraph import Paragraph
import re


def replace_placeholder_in_paragraph(paragraph: Paragraph, placeholder: str, value: str) -> None:
    """在段落中替换占位符，支持跨run的情况"""
    # 限制最大替换次数以防无限循环
    max_replacements = 50
    count = 0
    
    while placeholder in paragraph.text and count < max_replacements:
        count += 1
        text = paragraph.text
        start_index = text.find(placeholder)
        if start_index == -1:
            break
        end_index = start_index + len(placeholder)
        
        runs = paragraph.runs
        current_pos = 0
        start_run_idx = -1
        end_run_idx = -1
        start_run_char_idx = -1
        end_run_char_idx = -1
        
        # 定位占位符所在的 run
        for i, run in enumerate(runs):
            run_len = len(run.text)
            run_end = current_pos + run_len
            
            if start_run_idx == -1 and start_index < run_end:
                start_run_idx = i
                start_run_char_idx = start_index - current_pos
            
            if end_run_idx == -1 and end_index <= run_end:
                end_run_idx = i
                end_run_char_idx = end_index - current_pos
            
            current_pos += run_len
        
        # 执行替换
        if start_run_idx != -1 and end_run_idx != -1:
            if start_run_idx == end_run_idx:
                # 简单情况：在同一个 run 内
                run = runs[start_run_idx]
                run.text = run.text[:start_run_char_idx] + str(value) + run.text[end_run_char_idx:]
            else:
                # 复杂情况：跨多个 run
                # 1. 处理起始 Run：保留前缀，追加值
                start_run = runs[start_run_idx]
                start_run.text = start_run.text[:start_run_char_idx] + str(value)
                
                # 2. 处理中间 Runs：清空
                for i in range(start_run_idx + 1, end_run_idx):
                    runs[i].text = ""
                
                # 3. 处理结束 Run：保留后缀
                end_run = runs[end_run_idx]
                end_run.text = end_run.text[end_run_char_idx:]


def process_paragraph(paragraph: Paragraph, data: dict) -> None:
    """处理段落中的占位符，支持跨 run 的占位符"""
    # 快速检查：如果段落中没有 {{，则无需处理
    if "{{" not in paragraph.text:
        return

    # 遍历所有数据进行替换
    for key, value in data.items():
        placeholder = f"{{{{{key}}}}}"
        if placeholder in paragraph.text:
            replace_placeholder_in_paragraph(paragraph, placeholder, value)


def process_table(table: Table, data: dict) -> None:
    """处理表格中的占位符"""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                process_paragraph(paragraph, data)


def process_header_footer(header_footer: _Header | _Footer, data: dict) -> None:
    """处理页眉或页脚中的占位符"""
    for paragraph in header_footer.paragraphs:
        process_paragraph(paragraph, data)
    
    # 处理页眉页脚中的表格
    for table in header_footer.tables:
        process_table(table, data)


def process_document(doc: DocxDocument, data: dict) -> None:
    """处理整个文档"""
    # 处理普通段落
    for paragraph in doc.paragraphs:
        process_paragraph(paragraph, data)
    
    # 处理表格
    for table in doc.tables:
        process_table(table, data)
    
    # 处理页眉和页脚
    for section in doc.sections:
        # 处理页眉
        process_header_footer(section.header, data)
        
        # 处理页脚
        process_header_footer(section.footer, data)
        
        # 处理不同第一页的页眉页脚
        if section.different_first_page_header_footer:
            process_header_footer(section.first_page_header, data)
            process_header_footer(section.first_page_footer, data)


def extract_placeholders_from_text(text: str) -> list[str]:
    """从文本中提取{{field_name}}格式的占位符"""
    pattern = r'{{([^}]+)}}'
    matches = re.findall(pattern, text)
    return matches


def extract_process_paragraph(paragraph: Paragraph, placeholders: set[str]) -> None:
    """处理段落中的占位符"""
    # 从段落文本中提取占位符
    text_placeholders = extract_placeholders_from_text(paragraph.text)
    for placeholder in text_placeholders:
        placeholders.add(placeholder)
    
    # 处理段落中的所有运行元素
    for run in paragraph.runs:
        run_placeholders = extract_placeholders_from_text(run.text)
        for placeholder in run_placeholders:
            placeholders.add(placeholder)


def extract_process_table(table: Table, placeholders: set[str]) -> None:
    """处理表格中的占位符"""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                extract_process_paragraph(paragraph, placeholders)


def extract_process_header_footer(header_footer: _Header | _Footer, placeholders: set[str]) -> None:
    """处理页眉或页脚中的占位符"""
    for paragraph in header_footer.paragraphs:
        extract_process_paragraph(paragraph, placeholders)
    
    # 处理页眉页脚中的表格
    for table in header_footer.tables:
        extract_process_table(table, placeholders)


def extract_placeholders(doc: DocxDocument) -> list[str]:
    """从整个文档中提取占位符"""
    placeholders = set()
    
    # 处理普通段落
    for paragraph in doc.paragraphs:
        extract_process_paragraph(paragraph, placeholders)
    
    # 处理表格
    for table in doc.tables:
        extract_process_table(table, placeholders)
    
    # 处理页眉和页脚
    for section in doc.sections:
        # 处理页眉
        extract_process_header_footer(section.header, placeholders)
        
        # 处理页脚
        extract_process_header_footer(section.footer, placeholders)
        
        # 处理不同第一页的页眉页脚
        if section.different_first_page_header_footer:
            extract_process_header_footer(section.first_page_header, placeholders)
            extract_process_header_footer(section.first_page_footer, placeholders)
    
    # 返回排序后的占位符列表
    return sorted(list(placeholders))


def main():
    """主函数"""
    # 解析命令行参数
    if len(sys.argv) < 2:
        print("用法: python main.py <模板路径> [输出路径] [JSON数据] [操作类型]")
        print("操作类型: fill (默认) | extract")
        sys.exit(1)
    
    template_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else ""
    json_data = sys.argv[3] if len(sys.argv) > 3 else ""
    operation = sys.argv[4] if len(sys.argv) > 4 else "fill"
    
    try:
        # 加载模板
        doc = Document(template_path)
        
        if operation == "extract":
            # 提取占位符
            placeholders = extract_placeholders(doc)
            # 输出JSON格式的占位符列表
            print(json.dumps(placeholders, ensure_ascii=False))
            sys.exit(0)
        elif operation == "fill":
            # 填充模板
            if not output_path:
                print("填充操作需要指定输出路径")
                sys.exit(1)
            
            if not json_data:
                print("填充操作需要指定JSON数据")
                sys.exit(1)
            
            # 解析JSON数据
            data = json.loads(json_data)
            
            # 处理文档
            process_document(doc, data)
            
            # 保存新文档
            doc.save(output_path)
            
            print(f"成功生成文件: {output_path}")
            sys.exit(0)
        else:
            print(f"未知操作类型: {operation}")
            sys.exit(1)
    except json.JSONDecodeError as e:
        print(f"JSON解析错误: {e}")
        sys.exit(1)
    except Exception as e:
        print(f"处理错误: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
