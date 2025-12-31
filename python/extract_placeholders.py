#!/usr/bin/env python3
"""
提取Word模板中的占位符
支持普通段落、表格单元格、页眉和页脚中的{{field_name}}占位符
"""

import sys
import json
from docx import Document
from docx.document import Document as DocxDocument
from docx.section import _Header, _Footer
from docx.table import Table
from docx.text.paragraph import Paragraph
import re


def extract_placeholders_from_text(text: str) -> list[str]:
    """从文本中提取{{field_name}}格式的占位符"""
    pattern = r'{{([^}]+)}}'
    matches = re.findall(pattern, text)
    return matches


def process_paragraph(paragraph: Paragraph, placeholders: set[str]) -> None:
    """处理段落中的占位符"""
    # 从段落文本中提取占位符
    text_placeholders = extract_placeholders_from_text(paragraph.text)
    for placeholder in text_placeholders:
        placeholders.add(placeholder)
    
    # 处理段落中的所有运行元素
    for run in paragraph.runs:
        run_placeholders = extract_placeholders_from_text(run.text)
        for placeholder in run_placeholders:
            placeholders.add(placeholder)


def process_table(table: Table, placeholders: set[str]) -> None:
    """处理表格中的占位符"""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                process_paragraph(paragraph, placeholders)


def process_header_footer(header_footer: _Header | _Footer, placeholders: set[str]) -> None:
    """处理页眉或页脚中的占位符"""
    for paragraph in header_footer.paragraphs:
        process_paragraph(paragraph, placeholders)
    
    # 处理页眉页脚中的表格
    for table in header_footer.tables:
        process_table(table, placeholders)


def extract_placeholders(doc: DocxDocument) -> list[str]:
    """从整个文档中提取占位符"""
    placeholders = set()
    
    # 处理普通段落
    for paragraph in doc.paragraphs:
        process_paragraph(paragraph, placeholders)
    
    # 处理表格
    for table in doc.tables:
        process_table(table, placeholders)
    
    # 处理页眉和页脚
    for section in doc.sections:
        # 处理页眉
        process_header_footer(section.header, placeholders)
        
        # 处理页脚
        process_header_footer(section.footer, placeholders)
        
        # 处理不同第一页的页眉页脚
        if section.different_first_page_header_footer:
            process_header_footer(section.first_page_header, placeholders)
            process_header_footer(section.first_page_footer, placeholders)
    
    # 返回排序后的占位符列表
    return sorted(list(placeholders))


def main():
    """主函数"""
    if len(sys.argv) != 2:
        print("用法: python extract_placeholders.py <模板路径>")
        sys.exit(1)
    
    template_path = sys.argv[1]
    
    try:
        # 加载模板
        doc = Document(template_path)
        
        # 提取占位符
        placeholders = extract_placeholders(doc)
        
        # 输出JSON格式的占位符列表
        print(json.dumps(placeholders, ensure_ascii=False, indent=2))
        sys.exit(0)
    except Exception as e:
        print(f"处理错误: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
